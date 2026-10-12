from pathlib import Path
import csv
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent

SWAT_DIR = BASE_DIR / "data" / "swat"
OUTPUT_DIR = SWAT_DIR / "dashboard"

NORMAL_FILE = SWAT_DIR / "SWaT_Dataset_Normal_v1.xlsx"
ATTACK_FILE = SWAT_DIR / "SWaT_Dataset_Attack_v0.xlsx"

NORMAL_OUTPUT = OUTPUT_DIR / "normal_sample.csv"
ATTACK_OUTPUT = OUTPUT_DIR / "attack_sample.csv"

# Maximum number of rows to keep in each dashboard sample.
MAX_ROWS = 20000


def clean_value(value):
    if isinstance(value, str):
        return value.strip()
    return value


def prepare_dataset(input_file, output_file, max_rows=MAX_ROWS):
    if not input_file.exists():
        raise FileNotFoundError(f"Dataset not found: {input_file}")

    print(f"\nReading: {input_file.name}")

    workbook = load_workbook(
        input_file,
        read_only=True,
        data_only=True
    )

    worksheet = workbook.active

    try:
        rows = worksheet.iter_rows(values_only=True)

        # SWaT spreadsheets commonly contain metadata in row 1.
        # The column headers are in row 2.
        next(rows, None)
        raw_headers = next(rows, None)

        if not raw_headers:
            raise ValueError(f"No column headers found in {input_file.name}")

        headers = []
        for index, header in enumerate(raw_headers):
            if header is None or not str(header).strip():
                headers.append(f"column_{index}")
            else:
                headers.append(str(header).strip())

        label_index = next(
            (
                i for i, header in enumerate(headers)
                if header.lower() in {
                    "label", "attack", "attack_label", "normal/attack"
                }
            ),
            None
        )

        # Keep rows at a regular interval to create a smaller sample.
        total_data_rows = max(0, worksheet.max_row - 2)
        stride = max(1, total_data_rows // max_rows)

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        written = 0
        seen = 0

        with output_file.open("w", newline="", encoding="utf-8-sig") as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(headers)

            for row in rows:
                seen += 1

                if written >= max_rows:
                    break

                if (seen - 1) % stride != 0:
                    continue

                values = [
                    clean_value(value)
                    for value in row[:len(headers)]
                ]

                # Pad incomplete rows if necessary.
                values += [""] * (len(headers) - len(values))

                # Normalize inconsistent SWaT attack labels.
                if label_index is not None and label_index < len(values):
                    label = str(values[label_index]).strip().lower()

                    if label in {"a ttack", "attack"}:
                        values[label_index] = "Attack"
                    elif label == "normal":
                        values[label_index] = "Normal"

                writer.writerow(values)
                written += 1

        print(f"Saved: {output_file}")
        print(f"Rows written: {written}")
        print(f"Columns: {len(headers)}")

    finally:
        workbook.close()


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    prepare_dataset(NORMAL_FILE, NORMAL_OUTPUT)
    prepare_dataset(ATTACK_FILE, ATTACK_OUTPUT)

    print("\nSWaT dashboard samples prepared successfully!")
    print(f"Normal sample: {NORMAL_OUTPUT}")
    print(f"Attack sample: {ATTACK_OUTPUT}")


if __name__ == "__main__":
    main()